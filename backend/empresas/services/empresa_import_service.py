import io
import re
import openpyxl as opx
from openpyxl.worksheet.datavalidation import DataValidation

from empresas.models import Empresa

# Valor gravado no banco → rótulo amigável (igual ao select do frontend)
REGIME_OPCOES = {
    "SIMPLES_NACIONAL": "Simples Nacional",
    "LUCRO_PRESUMIDO": "Lucro Presumido",
    "LUCRO_REAL": "Lucro Real",
    "MEI": "MEI",
    "OUTROS": "Outros",
}

# Aceita código OU nome (com variações de escrita) → código oficial
REGIME_ALIASES = {
    "SIMPLES_NACIONAL": "SIMPLES_NACIONAL",
    "SIMPLES NACIONAL": "SIMPLES_NACIONAL",
    "LUCRO_PRESUMIDO": "LUCRO_PRESUMIDO",
    "LUCRO PRESUMIDO": "LUCRO_PRESUMIDO",
    "LUCRO_REAL": "LUCRO_REAL",
    "LUCRO REAL": "LUCRO_REAL",
    "MEI": "MEI",
    "OUTROS": "OUTROS",
}


def normalizar_regime(valor) -> str | None:
    """Converte o que veio do Excel para o código do REGIME_CHOICES."""
    if valor is None:
        return None

    texto = str(valor).strip()
    if not texto:
        return None

    chave = re.sub(r"\s+", " ", texto).upper()
    return REGIME_ALIASES.get(chave)


def importar_empresas_excel(arquivo_excel):
    workbook = opx.load_workbook(arquivo_excel)
    sheet = workbook.active

    empresas_importadas = []
    erros = []

    # Linha 1 = cabeçalho
    # Colunas: A=Razão Social | B=CNPJ | C=Regime
    for indice, linha in enumerate(sheet.iter_rows(min_row=2), start=2):
        razao_social = linha[0].value
        cnpj = linha[1].value
        regime_bruto = linha[2].value

        if not razao_social and not cnpj and not regime_bruto:
            continue

        if not razao_social or not cnpj or not regime_bruto:
            erros.append(f"Linha {indice}: preencha razão social, CNPJ e regime.")
            continue

        razao_social = str(razao_social).strip()

        cnpj_limpo = re.sub(r"\D", "", str(cnpj))
        if len(cnpj_limpo) != 14:
            erros.append(f"Linha {indice}: CNPJ inválido ({cnpj}).")
            continue

        regime = normalizar_regime(regime_bruto)
        if not regime:
            erros.append(
                f"Linha {indice}: regime inválido ({regime_bruto}). "
                "Use o menu da planilha modelo."
            )
            continue

        existente = Empresa.objects.filter(cnpj=cnpj_limpo).first()
        if existente is not None:
            if not existente.ativo:
                erros.append(
                    f"Linha {indice}: empresa existe porém está inativa "
                    f"(CNPJ {cnpj_limpo})."
                )
            else:
                erros.append(
                    f"Linha {indice}: empresa já existe (CNPJ {cnpj_limpo})."
                )
            continue

        empresa = Empresa.objects.create(
            razao_social=razao_social,
            cnpj=cnpj_limpo,
            regime=regime,
            ativo=True,
        )

        empresas_importadas.append({
            "id": empresa.id,
            "razao_social": empresa.razao_social,
            "cnpj": empresa.cnpj,
            "regime": empresa.regime,
            "criada": True,
        })

    return {
        "empresas": empresas_importadas,
        "erros": erros,
    }


def gerar_modelo_excel_empresas() -> bytes:
    """
    Gera .xlsx de exemplo com:
    - cabeçalho
    - 1 linha de exemplo
    - dropdown na coluna Regime (Data Validation do Excel)
    """
    workbook = opx.Workbook()
    sheet = workbook.active
    sheet.title = "Empresas"

    sheet["A1"] = "Razão Social"
    sheet["B1"] = "CNPJ"
    sheet["C1"] = "Regime"

    # Exemplo preenchido
    sheet["A2"] = "EMPRESA EXEMPLO LTDA"
    sheet["B2"] = "12345678000199"
    sheet["C2"] = "SIMPLES_NACIONAL"

    # Lista oculta com as opções oficiais (códigos do sistema)
    opcoes_sheet = workbook.create_sheet("opcoes_regime")
    for i, codigo in enumerate(REGIME_OPCOES.keys(), start=1):
        opcoes_sheet.cell(row=i, column=1, value=codigo)
    opcoes_sheet.sheet_state = "hidden"

    validacao = DataValidation(
        type="list",
        formula1="=opcoes_regime!$A$1:$A$5",
        allow_blank=False,
        showDropDown=False,  # False = mostra a setinha do menu
        showErrorMessage=True,
        errorTitle="Regime inválido",
        error="Selecione uma opção do menu (não digite livremente).",
        promptTitle="Regime",
        prompt="Escolha o regime tributário",
    )
    sheet.add_data_validation(validacao)
    validacao.add("C2:C1000")

    sheet.column_dimensions["A"].width = 40
    sheet.column_dimensions["B"].width = 20
    sheet.column_dimensions["C"].width = 22

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
